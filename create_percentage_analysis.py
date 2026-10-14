from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import DoughnutChart, BarChart, Reference
from openpyxl.utils import get_column_letter

ROOT=Path('/home/ubuntu/Simple-Percentage-Analysis')
ROOT.mkdir(parents=True, exist_ok=True)
CSV=Path('/home/ubuntu/work/superstore.csv')
OUT=ROOT/'simple_percentage_analysis.xlsx'

df=pd.read_csv(CSV)
summary=(df.groupby('Category',as_index=False).agg(Total_Sales=('Sales','sum'),Orders=('Order ID','nunique'),Quantity=('Quantity','sum'),Profit=('Profit','sum')).sort_values('Total_Sales',ascending=False))
total_sales=float(summary['Total_Sales'].sum())
summary['Contribution_Percentage']=summary['Total_Sales']/total_sales

wb=Workbook()
ws=wb.active; ws.title='Percentage_Analysis'
raw=wb.create_sheet('Source_Data')
notes=wb.create_sheet('Calculation_Notes')
navy='17365D'; blue='2F75B5'; light='D9EAF7'; pale='F4F8FB'; green='E2F0D9'; orange='FCE4D6'; white='FFFFFF'; gray='666666'; thin=Side(style='thin',color='D9E2F3')

def title(sheet,text,subtitle,end_col):
    sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=end_col)
    c=sheet.cell(1,1,text); c.font=Font(size=17,bold=True,color=white); c.fill=PatternFill('solid',fgColor=navy); c.alignment=Alignment(horizontal='left'); sheet.row_dimensions[1].height=30
    sheet.merge_cells(start_row=2,start_column=1,end_row=2,end_column=end_col)
    c=sheet.cell(2,1,subtitle); c.font=Font(size=10,italic=True,color=gray); c.alignment=Alignment(wrap_text=True)

def header(sheet,row,n):
    for col in range(1,n+1):
        c=sheet.cell(row,col); c.font=Font(bold=True,color=white); c.fill=PatternFill('solid',fgColor=blue); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=Border(bottom=thin)
    sheet.row_dimensions[row].height=26

def widths(sheet,n):
    for col in range(1,n+1):
        letter=get_column_letter(col); maxlen=max((len(str(c.value)) for c in sheet[letter] if c.value is not None),default=10)
        sheet.column_dimensions[letter].width=min(max(maxlen+2,14),32)

# Main table with formula-driven contribution percentage.
title(ws,'Simple Percentage Analysis','Sample Superstore sales contribution by category. Denominator: total sales across all categories.',7)
headers=['Category','Total Sales','Contribution %','Orders','Quantity','Profit','Check']
for j,h in enumerate(headers,1): ws.cell(4,j,h)
header(ws,4,len(headers))
for i,row in enumerate(summary.itertuples(index=False),5):
    ws.cell(i,1,row.Category)
    ws.cell(i,2,float(row.Total_Sales))
    ws.cell(i,3,f'=B{i}/$B${5+len(summary)}')
    ws.cell(i,4,int(row.Orders)); ws.cell(i,5,int(row.Quantity)); ws.cell(i,6,float(row.Profit)); ws.cell(i,7,f'=C{i}')
total_row=5+len(summary)
ws.cell(total_row,1,'TOTAL'); ws.cell(total_row,2,f'=SUM(B5:B{total_row-1})'); ws.cell(total_row,3,f'=SUM(C5:C{total_row-1})'); ws.cell(total_row,4,f'=SUM(D5:D{total_row-1})'); ws.cell(total_row,5,f'=SUM(E5:E{total_row-1})'); ws.cell(total_row,6,f'=SUM(F5:F{total_row-1})'); ws.cell(total_row,7,f'=C{total_row}')
for c in range(1,8): ws.cell(total_row,c).font=Font(bold=True,color=navy); ws.cell(total_row,c).fill=PatternFill('solid',fgColor=green)
ref=f'A4:G{total_row-1}'; tab=Table(displayName='CategoryPercentage',ref=ref); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showColumnStripes=False); ws.add_table(tab)
for r in range(5,total_row+1):
    ws.cell(r,2).number_format='$#,##0.00'; ws.cell(r,3).number_format='0.00%'; ws.cell(r,6).number_format='$#,##0.00'; ws.cell(r,7).number_format='0.00%'
ws.freeze_panes='A5'; widths(ws,7)
# Charts
chart=DoughnutChart(); chart.title='Sales Contribution by Category'; chart.holeSize=55; chart.height=8; chart.width=13
chart.add_data(Reference(ws,min_col=2,min_row=4,max_row=total_row-1),titles_from_data=True); chart.set_categories(Reference(ws,min_col=1,min_row=5,max_row=total_row-1)); ws.add_chart(chart,'I4')
bar=BarChart(); bar.type='bar'; bar.style=10; bar.title='Total Sales by Category'; bar.y_axis.title='Category'; bar.x_axis.title='Sales'; bar.height=8; bar.width=13
bar.add_data(Reference(ws,min_col=2,min_row=4,max_row=total_row-1),titles_from_data=True); bar.set_categories(Reference(ws,min_col=1,min_row=5,max_row=total_row-1)); ws.add_chart(bar,'I20')

# Source data sheet with original dataset.
title(raw,'Sample Superstore Source Data','Source records used to calculate category sales percentages.',len(df.columns))
for j,c in enumerate(df.columns,1): raw.cell(4,j,c)
header(raw,4,len(df.columns))
for i,row in enumerate(df.itertuples(index=False,name=None),5):
    for j,v in enumerate(row,1): raw.cell(i,j,v)
ref=f'A4:{get_column_letter(len(df.columns))}{4+len(df)}'; tab=Table(displayName='SuperstoreSource',ref=ref); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showColumnStripes=False); raw.add_table(tab)
raw.freeze_panes='A5'; widths(raw,len(df.columns))

# Notes sheet explains formula and denominator.
title(notes,'Calculation Notes','Method used for the percentage analysis.',4)
items=[('Dataset','Sample Superstore sales dataset'),('Rows analyzed',len(df)),('Grouping field','Category'),('Measure','Sales'),('Denominator','Total sales across Furniture, Office Supplies, and Technology'),('Formula','Category Sales ÷ Total Sales × 100'),('Validation','Contribution percentages sum to 100.00%'),('Source','https://github.com/leonism/sample-superstore')]
for i,(k,v) in enumerate(items,4):
    notes.cell(i,1,k).font=Font(bold=True,color=navy); notes.cell(i,1).fill=PatternFill('solid',fgColor=light); notes.cell(i,2,v).fill=PatternFill('solid',fgColor=pale); notes.cell(i,2).alignment=Alignment(wrap_text=True)
notes['A14']='Interview notes'; notes['A14'].font=Font(size=13,bold=True,color=navy)
notes['A15']='Percentage contribution shows how much each category contributes to the selected total. The denominator must be stated so the result is clear and reproducible.'; notes['A15'].alignment=Alignment(wrap_text=True)
notes.column_dimensions['A'].width=24; notes.column_dimensions['B'].width=88
for sheet in wb.worksheets: sheet.sheet_view.showGridLines=False
wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True
wb.save(OUT)
print(OUT)
print('records',len(df),'categories',len(summary),'total_sales',round(total_sales,2),'percent_total',round(summary['Contribution_Percentage'].sum()*100,2))
